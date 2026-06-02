# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""
Cross-check data across all pipeline stages to ensure completeness.

Verifies:
  Step 1: Every preparer folder has a matching assessor evaluation (and vice versa)
  Step 2: Every assessment appears in the analyst reports (markdown + Excel)

Usage:
    uv run .github/skills/evaluation-summarization/scripts/verify_coverage.py \
        --preparer-dir data/preparer/ --assessor-dir data/assessor/ --analyst-dir data/analyst/
"""

import argparse
import io
import json
import re
import sys
from datetime import datetime
from pathlib import Path

# Force UTF-8 output on Windows
if sys.stdout.encoding != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def is_skeleton(e: dict) -> bool:
    hook = e.get("candidate_hook", "")
    skeleton_phrases = ("No submission provided", "No responses provided", "Empty file", "Empty submission")
    return any(phrase in hook for phrase in skeleton_phrases)


SCORE_ALIASES = {
    "written": ["written_communications", "q1_written_communications", "Q1_written", "q1_written"],
    "oral": ["oral_reflective", "oral_communications", "oral_communications_reflective",
             "q2_oral_reflective", "Q2_oral", "q2_oral"],
    "logical": ["oral_logical", "logical_thinking", "oral_communications_logical",
                "q3_oral_logical", "q3_logical_thinking", "q3_oral_logical_thinking",
                "Q3_logical", "q3_logical"],
}


def get_score(e: dict, key: str):
    candidates = SCORE_ALIASES.get(key, [key])
    for wrapper in [e, e.get("assessments", {}), e.get("assessment_scores", {}), e.get("scores", {})]:
        if not isinstance(wrapper, dict):
            continue
        for c in candidates:
            if c in wrapper:
                sec = wrapper[c]
                if isinstance(sec, dict):
                    val = sec.get("final_score", sec.get("overall_score"))
                    if val is not None:
                        return val
    return None


def load_assessor_evaluations(assessor_dir: Path) -> dict[str, dict]:
    """Return {applicant_id: data} for all evaluation.json files in assessor subdirs."""
    result = {}
    for applicant_dir in sorted(assessor_dir.iterdir()):
        if not applicant_dir.is_dir():
            continue
        eval_file = applicant_dir / "evaluation.json"
        if not eval_file.exists():
            continue
        try:
            with open(eval_file, encoding="utf-8") as f:
                data = json.load(f)
        except json.JSONDecodeError:
            continue
        aid = applicant_dir.name
        if not data.get("applicant_id") or data["applicant_id"] == "unknown":
            data["applicant_id"] = aid
        result[aid] = data
    return result


# ---------------------------------------------------------------------------
# Step 1: Preparer ↔ Assessor mapping
# ---------------------------------------------------------------------------

def step1_mapping(preparer_ids: set[str], assessor_ids: set[str]) -> tuple[bool, list[str]]:
    lines = []
    missing_evals = sorted(preparer_ids - assessor_ids)
    orphan_evals = sorted(assessor_ids - preparer_ids)

    lines.append("STEP 1: Preparer ↔ Assessor mapping")
    lines.append(f"  Preparer folders     : {len(preparer_ids)}")
    lines.append(f"  Assessor evaluations : {len(assessor_ids)}")
    lines.append(f"  Missing evaluations  : {len(missing_evals)}")
    lines.append(f"  Orphan evaluations   : {len(orphan_evals)}")

    if missing_evals:
        lines.append("  Preparer folders with no assessor evaluation:")
        for aid in missing_evals:
            lines.append(f"    - {aid}")
    if orphan_evals:
        lines.append("  Assessor evaluations with no preparer folder:")
        for aid in orphan_evals:
            lines.append(f"    - {aid}")

    passed = not missing_evals and not orphan_evals
    lines.append(f"  Status: {'PASS' if passed else 'FAIL'}")
    return passed, lines


# ---------------------------------------------------------------------------
# Step 2: Assessor → Analyst report coverage
# ---------------------------------------------------------------------------

def step2_reports(assessor_map: dict[str, dict], analyst_dir: Path) -> tuple[bool, list[str]]:
    lines = []

    real_ids = {aid for aid, data in assessor_map.items() if not is_skeleton(data)}

    # Read markdown headings from aggregate report
    md_path = analyst_dir / "evaluation_report.md"
    md_ids: set[str] = set()
    if md_path.exists():
        text = md_path.read_text(encoding="utf-8")
        md_ids = set(re.findall(r"^### Applicant (\S+)", text, re.MULTILINE))

    # Read xlsx Scores sheet column A (skip header)
    xlsx_path = analyst_dir / "evaluation_summary.xlsx"
    xlsx_ids: set[str] = set()
    if xlsx_path.exists():
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb["Scores"]
        for i, row in enumerate(ws.iter_rows(min_col=1, max_col=1, values_only=True)):
            if i == 0:
                continue
            val = row[0]
            if val:
                xlsx_ids.add(str(val).strip())
        wb.close()

    missing_md = sorted(real_ids - md_ids)
    missing_xlsx = sorted(real_ids - xlsx_ids)

    lines.append("STEP 2: Report coverage")
    lines.append(f"  Non-skeleton evaluations : {len(real_ids)}")
    lines.append(f"  Found in .md             : {len(real_ids & md_ids)} / {len(real_ids)}")
    lines.append(f"  Found in .xlsx           : {len(real_ids & xlsx_ids)} / {len(real_ids)}")
    lines.append(f"  MISSING_MD               : {len(missing_md)}")
    lines.append(f"  MISSING_XLSX             : {len(missing_xlsx)}")

    if missing_md:
        lines.append("  Evaluations absent from evaluation_report.md:")
        for aid in missing_md:
            lines.append(f"    - {aid}")
    if missing_xlsx:
        lines.append("  Evaluations absent from evaluation_summary.xlsx:")
        for aid in missing_xlsx:
            lines.append(f"    - {aid}")

    passed = not missing_md and not missing_xlsx
    lines.append(f"  Status: {'PASS' if passed else 'FAIL'}")
    return passed, lines


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Verify data coverage across pipeline stages"
    )
    p.add_argument(
        "--preparer-dir", required=True, type=Path,
        help="Directory with preparer applicant folders (e.g. data/preparer/)"
    )
    p.add_argument(
        "--assessor-dir", required=True, type=Path,
        help="Directory with assessor evaluation folders (e.g. data/assessor/)"
    )
    p.add_argument(
        "--analyst-dir", required=True, type=Path,
        help="Directory with analyst output files (e.g. data/analyst/)"
    )
    return p.parse_args()


def main() -> int:
    args = parse_args()
    preparer_dir = args.preparer_dir.resolve()
    assessor_dir = args.assessor_dir.resolve()
    analyst_dir = args.analyst_dir.resolve()

    # Validate directories exist
    for label, d in [("preparer", preparer_dir), ("assessor", assessor_dir)]:
        if not d.is_dir():
            log(f"ERROR: {label} directory does not exist: {d}")
            return 1

    # Collect IDs
    preparer_ids = {p.name for p in preparer_dir.iterdir() if p.is_dir() and not p.name.startswith(".")}
    assessor_map = load_assessor_evaluations(assessor_dir)
    assessor_ids = set(assessor_map.keys())

    header = [
        "=== BBA Evaluation Coverage Report ===",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
    ]

    p1, l1 = step1_mapping(preparer_ids, assessor_ids)

    # Step 2 only if analyst dir exists
    if analyst_dir.is_dir():
        p2, l2 = step2_reports(assessor_map, analyst_dir)
    else:
        p2 = True
        l2 = ["STEP 2: Report coverage", "  Analyst directory not found — skipping report checks", "  Status: SKIP"]

    all_pass = p1 and p2
    footer = [
        "",
        f"OVERALL: {'ALL CHECKS PASSED' if all_pass else 'FAILURES DETECTED — see details above'}",
    ]

    report_lines = header + l1 + [""] + l2 + footer
    report = "\n".join(report_lines)

    # Write verification report
    analyst_dir.mkdir(parents=True, exist_ok=True)
    output_path = analyst_dir / "verification_report.txt"
    output_path.write_text(report, encoding="utf-8")
    log(f"Written: {output_path}")

    # JSON summary to stdout
    summary = {
        "action": "verify_coverage",
        "preparer_dir": str(preparer_dir),
        "assessor_dir": str(assessor_dir),
        "analyst_dir": str(analyst_dir),
        "preparer_count": len(preparer_ids),
        "assessor_count": len(assessor_ids),
        "step1_passed": p1,
        "step2_passed": p2,
        "all_passed": all_pass,
        "report_file": str(output_path),
    }
    print(json.dumps(summary, indent=2))

    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
