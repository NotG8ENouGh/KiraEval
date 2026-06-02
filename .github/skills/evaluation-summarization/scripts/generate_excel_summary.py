# /// script
# requires-python = ">=3.11"
# dependencies = ["pyyaml>=6.0", "openpyxl>=3.1"]
# ///
"""
Generate Excel summary spreadsheet from individual evaluation JSONs.

Reads per-question assessment files from per-applicant assessor directories and produces
an Excel workbook with scores and detailed statistics.

New schema (preferred): reads q1_written.json + q2_oral.json + q3_logical.json per applicant.
Legacy schema (fallback): reads evaluation.json.

Usage:
    uv run .github/skills/evaluation-summarization/scripts/generate_excel_summary.py \
        --input-dir data/assessor/ --output-dir data/analyst/ [--preparer-dir data/preparer/]
"""

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from statistics import mean, median, stdev, variance

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


# ---------------------------------------------------------------------------
# Score extraction helpers
# ---------------------------------------------------------------------------

SCORE_ALIASES = {
    "written": ["written_communications", "q1_written_communications", "Q1_written", "q1_written"],
    "oral": ["oral_reflective", "oral_communications", "oral_communications_reflective",
             "q2_oral_reflective", "Q2_oral", "q2_oral"],
    "logical": ["oral_logical", "logical_thinking", "oral_communications_logical",
                "q3_oral_logical", "q3_logical_thinking", "q3_oral_logical_thinking",
                "Q3_logical", "q3_logical"],
}


def get_score(e: dict, key: str) -> float | None:
    candidates = SCORE_ALIASES.get(key, [key])
    for wrapper in [e, e.get("assessments", {}), e.get("assessment_scores", {}), e.get("scores", {})]:
        for c in candidates:
            if c in wrapper:
                sec = wrapper[c]
                val = sec.get("final_score", sec.get("overall_score"))
                if val is not None:
                    return val
    return None


def is_skeleton(e: dict) -> bool:
    hook = e.get("candidate_hook", "")
    skeleton_phrases = ("No submission provided", "No responses provided", "Empty file", "Empty submission")
    return any(phrase in hook for phrase in skeleton_phrases)


# Maps per-question JSON filename to its section key
_Q_FILE_SECTION_MAP = [
    ("q1_written.json", "written_communications"),
    ("q2_oral.json", "oral_communications"),
    ("q3_logical.json", "logical_thinking"),
]


def load_evaluation_for_applicant(applicant_dir: Path) -> dict | None:
    """Load from 3 per-question files (new schema) or fall back to evaluation.json."""
    has_new = any((applicant_dir / fname).exists() for fname, _ in _Q_FILE_SECTION_MAP)

    if has_new:
        merged: dict = {"applicant_id": applicant_dir.name}
        for fname, section_key in _Q_FILE_SECTION_MAP:
            json_file = applicant_dir / fname
            if not json_file.exists():
                continue
            try:
                data = json.loads(json_file.read_text(encoding="utf-8"))
            except json.JSONDecodeError as exc:
                log(f"WARNING: Skipping {json_file}: Invalid JSON — {exc}")
                continue
            if section_key in data:
                merged[section_key] = data[section_key]
            if data.get("applicant_id") and data["applicant_id"] != "unknown":
                merged["applicant_id"] = data["applicant_id"]
        return merged

    eval_file = applicant_dir / "evaluation.json"
    if not eval_file.exists():
        return None
    try:
        return json.loads(eval_file.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        log(f"WARNING: Skipping {applicant_dir.name}: Invalid JSON — {exc}")
        return None


def load_evaluations(input_dir: Path) -> list[dict]:
    evaluations = []
    for applicant_dir in sorted(input_dir.iterdir()):
        if not applicant_dir.is_dir():
            continue
        data = load_evaluation_for_applicant(applicant_dir)
        if data is None:
            continue
        if is_skeleton(data):
            log(f"WARNING: Skipping {applicant_dir.name}: empty placeholder")
            continue
        if not data.get("applicant_id") or data["applicant_id"] == "unknown":
            data["applicant_id"] = applicant_dir.name
        evaluations.append(data)
    return evaluations


def load_kira_urls(preparer_dir: Path | None, applicant_ids: list[str]) -> dict[str, str]:
    """Return {applicant_id: kira_url} for each applicant that has a meta.json with a URL."""
    if preparer_dir is None or not preparer_dir.is_dir():
        return {}
    urls: dict[str, str] = {}
    for aid in applicant_ids:
        meta_path = preparer_dir / aid / "meta.json"
        if not meta_path.exists():
            continue
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            url = meta.get("url")
            if url:
                urls[aid] = url
        except Exception:
            pass
    return urls


def calculate_distribution(scores: list[float]) -> dict[int, int]:
    rounded_scores = [round(score) for score in scores]
    distribution = Counter(rounded_scores)
    return {i: distribution.get(i, 0) for i in range(1, 6)}


def find_best_worst(evaluations: list[dict], assessment_type: str) -> dict:
    scores_map: dict[float, list[str]] = {}
    for e in evaluations:
        score = get_score(e, assessment_type)
        if score is None:
            continue
        scores_map.setdefault(score, []).append(e["applicant_id"])
    if not scores_map:
        return {"best": {"score": 0, "candidates": []}, "worst": {"score": 0, "candidates": []}}
    return {
        "best": {"score": max(scores_map), "candidates": sorted(scores_map[max(scores_map)])},
        "worst": {"score": min(scores_map), "candidates": sorted(scores_map[min(scores_map)])},
    }


# ---------------------------------------------------------------------------
# Excel generation
# ---------------------------------------------------------------------------

def add_scores_sheet(wb: Workbook, evaluations: list[dict], kira_urls: dict[str, str] | None = None) -> None:
    ws = wb.active
    ws.title = "Scores"

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    link_font = Font(color="0563C1", underline="single")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    headers = ["Candidate ID", "Written Score", "Oral Score", "Logical Score"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    kira_urls = kira_urls or {}

    def composite(e: dict) -> float:
        w = get_score(e, "written") or 0.0
        o = get_score(e, "oral") or 0.0
        l = get_score(e, "logical") or 0.0
        return w + o + l

    for row, evaluation in enumerate(sorted(evaluations, key=composite, reverse=True), 2):
        aid = evaluation["applicant_id"]
        w = get_score(evaluation, "written")
        o = get_score(evaluation, "oral")
        l = get_score(evaluation, "logical")

        cell = ws.cell(row=row, column=1, value=aid)
        cell.alignment = center_align
        cell.border = border
        url = kira_urls.get(aid)
        if url:
            cell.hyperlink = url
            cell.font = link_font

        for col, val in [(2, w), (3, o), (4, l)]:
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = center_align
            cell.border = border
            if val is not None:
                cell.number_format = "0.0"

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15


def add_statistics_sheet(wb: Workbook, evaluations: list[dict]) -> None:
    ws = wb.create_sheet("Statistics")

    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    written_scores, oral_scores, logical_scores = [], [], []
    for e in evaluations:
        w = get_score(e, "written")
        o = get_score(e, "oral")
        l = get_score(e, "logical")
        if w is not None:
            written_scores.append(w)
        if o is not None:
            oral_scores.append(o)
        if l is not None:
            logical_scores.append(l)

    def calc_stats(scores):
        if not scores:
            return {"mean": 0, "median": 0, "stdev": 0, "variance": 0, "min": 0, "max": 0, "n": 0}
        return {
            "mean": mean(scores), "median": median(scores),
            "stdev": stdev(scores) if len(scores) > 1 else 0,
            "variance": variance(scores) if len(scores) > 1 else 0,
            "min": min(scores), "max": max(scores),
            "n": len(scores),
        }

    written_stats = calc_stats(written_scores)
    oral_stats = calc_stats(oral_scores)
    logical_stats = calc_stats(logical_scores)

    # Title
    ws.merge_cells("A1:G1")
    cell = ws["A1"]
    cell.value = "Detailed Statistics"
    cell.font = title_font
    cell.alignment = center_align

    # Assessment Scores Statistics
    row = 3
    ws.cell(row=row, column=1, value="Assessment Scores").font = header_font
    row += 1
    for col, header in enumerate(["Assessment Type", "n", "Mean", "Median", "Std Dev", "Min", "Max"], 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    for row_idx, (label, s) in enumerate([
        ("Written Communications", written_stats),
        ("Oral Communications", oral_stats),
        ("Logical Thinking", logical_stats),
    ], row + 1):
        ws.cell(row=row_idx, column=1, value=label).border = border
        for col, (val, fmt) in enumerate([
            (s["n"], "0"), (s["mean"], "0.00"), (s["median"], "0.00"),
            (s["stdev"], "0.00"), (s["min"], "0.0"), (s["max"], "0.0"),
        ], 2):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.number_format = fmt
            c.border = border
            c.alignment = center_align

    # Score Distribution
    row = row_idx + 3
    ws.cell(row=row, column=1, value="Score Distribution").font = title_font

    for label, scores in [
        ("Written Communications", written_scores),
        ("Oral Communications", oral_scores),
        ("Logical Thinking", logical_scores),
    ]:
        row += 2
        ws.cell(row=row, column=1, value=label).font = header_font
        row += 1
        for col, header in enumerate(["Score", "Count", "Percentage"], 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border

        dist = calculate_distribution(scores)
        total = len(scores)
        for score in range(5, 0, -1):
            row += 1
            count = dist[score]
            pct = (count / total * 100) if total > 0 else 0
            ws.cell(row=row, column=1, value=score).border = border
            ws.cell(row=row, column=1).alignment = center_align
            ws.cell(row=row, column=2, value=count).border = border
            ws.cell(row=row, column=2).alignment = center_align
            c = ws.cell(row=row, column=3, value=pct)
            c.number_format = "0.0"
            c.border = border
            c.alignment = center_align

    # Performance Analysis
    row += 3
    ws.cell(row=row, column=1, value="Performance Analysis").font = title_font

    for assessment_type, label in [
        ("written", "Written Communications"),
        ("oral", "Oral Communications"),
        ("logical", "Logical Thinking"),
    ]:
        row += 2
        ws.cell(row=row, column=1, value=label).font = header_font
        row += 1
        for col, header in enumerate(["Category", "Score", "Candidates"], 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border

        bw = find_best_worst(evaluations, assessment_type)
        for label2, entry in [("Best", bw["best"]), ("Worst", bw["worst"])]:
            row += 1
            ws.cell(row=row, column=1, value=label2).border = border
            c = ws.cell(row=row, column=2, value=entry["score"])
            c.number_format = "0.0"
            c.border = border
            c.alignment = center_align
            ws.cell(row=row, column=3, value=", ".join(entry["candidates"])).border = border

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Generate Excel summary from assessor evaluation data"
    )
    p.add_argument(
        "--input-dir", required=True, type=Path,
        help="Directory containing per-applicant evaluation folders (e.g. data/assessor/)"
    )
    p.add_argument(
        "--output-dir", required=True, type=Path,
        help="Directory to write Excel file to (e.g. data/analyst/)"
    )
    p.add_argument(
        "--preparer-dir", required=False, type=Path, default=None,
        help="Optional preparer directory to read Kira Review URLs from meta.json (e.g. data/preparer/)"
    )
    return p.parse_args()


def main() -> int:
    args = parse_args()
    input_dir = args.input_dir.resolve()
    output_dir = args.output_dir.resolve()
    preparer_dir = args.preparer_dir.resolve() if args.preparer_dir else None

    if not input_dir.is_dir():
        log(f"ERROR: Input directory does not exist: {input_dir}")
        return 2

    log("Loading evaluations...")
    evaluations = load_evaluations(input_dir)

    if not evaluations:
        log("ERROR: No evaluation files found in input directory")
        return 2

    log(f"Loaded {len(evaluations)} evaluations")

    output_dir.mkdir(parents=True, exist_ok=True)

    # Load Kira URLs for hyperlinks
    kira_urls = load_kira_urls(preparer_dir, [e["applicant_id"] for e in evaluations])
    if kira_urls:
        log(f"Loaded Kira URLs for {len(kira_urls)} applicants")

    log("Creating Excel summary...")
    wb = Workbook()
    add_scores_sheet(wb, evaluations, kira_urls)
    add_statistics_sheet(wb, evaluations)

    output_path = output_dir / "evaluation_summary.xlsx"
    wb.save(output_path)
    log(f"Written: {output_path}")

    summary = {
        "action": "generate_excel_summary",
        "input_dir": str(input_dir),
        "output_dir": str(output_dir),
        "total_evaluations": len(evaluations),
        "output_file": str(output_path),
        "kira_urls_loaded": len(kira_urls),
    }
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
